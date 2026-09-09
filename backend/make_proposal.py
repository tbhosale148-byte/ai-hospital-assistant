from docx import Document

def create_safeskill_proposal():
    doc = Document()
    
    # Title Section
    doc.add_heading("SAFESKILL AI: PROJECT PROPOSAL & SPECIFICATIONS", level=0)
    doc.add_paragraph("Project Title: SafeSkill AI\nDomain: Intelligent Worker Safety & Performance Monitoring System for Manufacturing")
    
    # Part 1: Project Proposal Summary
    doc.add_heading("PART 1: PROJECT PROPOSAL", level=1)
    doc.add_heading("1. Executive Summary & Problem Statement", level=2)
    doc.add_paragraph("SafeSkill AI is an artificial intelligence capstone project designed to connect camera safety monitoring with work production data, helping factory managers prevent workplace accidents and product quality defects[cite: 1].\n\n"
                      "Core problems addressed include:\n"
                      "• Unreported Safety Hazards: Over 80% of safety near-misses (e.g., walking too close to forklift paths without safety gear) go unreported[cite: 1].\n"
                      "• Undetected Skill Drift: Workers slowly losing efficiency over time, resulting in higher product error rates discovered too late[cite: 1].")
    
    doc.add_heading("2. Key Deliverables & Objectives", level=2)
    doc.add_paragraph("• Computer Vision Safety Pipeline: Video detection system to spot workers, track forklifts, and catch near-miss events[cite: 1].\n"
                      "• Performance Risk Engine: ML model to predict worker error risks and skill drops based on daily logs[cite: 1].\n"
                      "• AI Report Generator: Automated assistant combining safety and work records into clear advisory reports[cite: 1].\n"
                      "• SafeSkill AI Web Portal: Dashboard providing live safety alerts, risk maps, and worker profiles[cite: 1].")

    # Part 2: SRS
    doc.add_heading("PART 2: SOFTWARE REQUIREMENTS SPECIFICATION (SRS)", level=1)
    doc.add_heading("1. Scope & User Classes", level=2)
    doc.add_paragraph("SafeSkill AI centralizes factory camera video streams and production database records to serve three main user classes[cite: 4]:\n"
                      "• Safety Manager: Needs live safety warnings, near-accident clips, and facility heatmaps[cite: 4].\n"
                      "• Plant Operations Manager: Needs worker productivity charts, error trends, and risk reports[cite: 4].\n"
                      "• HR/Training Specialist: Focuses on AI-generated personal training plans[cite: 4].")
    
    doc.add_heading("2. Functional & Non-Functional Requirements", level=2)
    doc.add_paragraph("• FR-CV-01 (Live Safety Detection): AI analysis for workers, machines, and restricted danger zones with immediate alerts[cite: 4].\n"
                      "• FR-ML-02 (Skill & Performance Risk Prediction): Generates overall Risk Scores from 0.0 to 1.0[cite: 4].\n"
                      "• NFR-PERF-01 (Processing Speed): Video detection must process at >= 15 frames per second[cite: 4].\n"
                      "• NFR-SEC-03 (Security): Strict role-based access control (RBAC) protecting all records[cite: 4].")

    # Part 3: TDS
    doc.add_heading("PART 3: TECHNICAL DESIGN SPECIFICATION (TDS)", level=1)
    doc.add_heading("1. System Architecture & Component Stack", level=2)
    doc.add_paragraph("• Layer 1 (Data Ingestion): Security cameras, video files, and production data uploads[cite: 3].\n"
                      "• Layer 2 (Analytics Engine): YOLOv8 vision processing and machine learning risk models[cite: 3].\n"
                      "• Layer 3 (Orchestration): AI assistant tools for generating advice[cite: 3].\n"
                      "• Layer 4 (Application Server): Python Flask web server and API connections[cite: 3].\n"
                      "• Layer 5 (Storage): Database records and media file storage[cite: 3].")
    
    doc.add_heading("2. Core Database Schema", level=2)
    doc.add_paragraph("• workers: Master employee directory (worker_id, name, department)[cite: 3].\n"
                      "• safety_incidents: Logged camera safety violations and severity[cite: 3].\n"
                      "• performance_logs: Daily work output, defect counts, and risk levels[cite: 3].\n"
                      "• ai_recommendations: Generated training plans and action items[cite: 3].")

    # Part 4: FSD
    doc.add_heading("PART 4: FUNCTIONAL DESIGN SPECIFICATION (FSD)", level=1)
    doc.add_heading("1. Key Functional Requirements & APIs", level=2)
    doc.add_paragraph("• FR-01 (Feed Ingestion): Reads video stream frame by frame[cite: 2].\n"
                      "• FR-02 (Zone Drawing): Custom safety zone outlines via screen coordinates[cite: 2].\n"
                      "• POST /api/v1/safety/analyze: Evaluates video stream and danger zones to return violations[cite: 3].\n"
                      "• POST /api/v1/performance/predict: Accepts worker metrics to return risk scores[cite: 3].\n"
                      "• POST /api/v1/reports/generate: Combines records to return AI action plans[cite: 3].")

    # Save document
    filename = "SafeSkillAI_Project_Proposal.docx"
    doc.save(filename)
    print(f"Success! Generated your file: {filename}")

if __name__ == "__main__":
    create_safeskill_proposal()